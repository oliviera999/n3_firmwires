"""Construit le classeur d'achats n3-universal (4 projets + global + prix) — appelé par build_achats.py."""
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).parent))
import inventaire as I
import catalogue as K
from prix import PRIX

from prix import DATE

F = "Arial"
f_base = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True, color="1F3A5F")
f_sub = Font(name=F, size=10, italic=True, color="555555")
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_input = Font(name=F, size=10, color="0000FF")
f_link = Font(name=F, size=10, color="1F5FBF", underline="single")
f_link_sheet = Font(name=F, size=10, color="008000")
f_warn = Font(name=F, size=10, color="9C2A00")
fill_head = PatternFill("solid", fgColor="1F3A5F")
fill_tot = PatternFill("solid", fgColor="E8EEF5")
fill_cat = [PatternFill("solid", fgColor=c) for c in ("FFFFFF", "F4F6F9")]
fill_opt = PatternFill("solid", fgColor="FFF6E0")
fill_own = PatternFill("solid", fgColor="EAF5EA")
fill_imp = PatternFill("solid", fgColor="FBE9E7")
thin = Side(style="thin", color="C9D1DC")
border = Border(bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
top = Alignment(vertical="top")
MONEY = '#,##0.00;-#,##0.00;"-"'

A, O, P = I.A, I.O, I.P

# ------------------------------------------------------------------ remarques d'usage carte
USE = {
    "ESP32_DEVKIT": "Enfiché sur ses supports, jamais soudé. Vérifier l'écart des rangées (25,4 ou 27,94 mm : support droit sur A ou A').",
    "FIL": "D5 = pont de fil (rev 0.2.2). Règle absolue : ne JAMAIS brancher l'USB du DevKit quand la carte est alimentée.",
    "LD1117": "Ce n'est pas un 78xx : 1=GND 2=OUT 3=IN (sérigraphié).",
    "IRF4905": "Respecter la sérigraphie G-D-S de chaque MOSFET.",
    "BC337": "Brochage C-B-E sérigraphié.",
    "HLK20M05": "Zone 230 V : montage et raccordement sous supervision adulte.",
    "TVS_6V8": "Protège le rail 5 V. Commander tôt chez LCSC (stock faible).",
    "VIS_M3": "",
}
CHECK_PINS = {"BARRETTE_F": "broches femelles utilisées", "BARRETTE_M": "broches mâles utilisées"}


def roles_text(lst, n=4):
    s = "; ".join(lst[:n])
    return s + (" ; …" if len(lst) > n else "")


def profile_rows(prof):
    """Liste ordonnée de dicts : id, cat, article, refs, role, qty, statut, remarque."""
    rows = []
    board = I.board_items(prof)
    notes = I.POSE_NOTE.get(prof, {})
    for aid, e in board.items():
        if e["qty"] == 0:
            continue
        rem = USE.get(aid, "")
        if aid in CHECK_PINS:
            rem = f"{e['pins']} {CHECK_PINS[aid]} (+ pertes de coupe)." + (" " + rem if rem else "")
        if aid == "CAVALIER":
            rem = notes.get("JP", "")
        if aid == "VIS_M3" and "H" in notes:
            rem = notes["H"]
        if aid == "BORNIER_2P":
            rem = " ".join(x for x in (notes.get("J1", ""), notes.get("J25", "")) if x)
        if aid == "R_2k":
            rem = "R17-R19 : ponts écho HC-SR04, uniquement sur ffp5cs."
        if aid == "LED_V":
            rem = "LED5 posée (présence 5 V)."
        refs = ", ".join(e["refs"]) if len(e["refs"]) <= 14 else ", ".join(e["refs"][:14]) + ", …"
        rows.append(dict(id=aid, refs=refs, role=roles_text(e["roles"]), qty=e["qty"], statut=A,
                         remarque=rem))
    for aid, q, st, role, rem in I.OFF[prof]:
        rows.append(dict(id=aid, refs="hors carte", role=role, qty=q, statut=st, remarque=rem))
    for r in rows:
        ci, des, _ = K.CAT[r["id"]]
        r["cat"], r["article"], r["ci"] = K.CATS[ci], des, ci
    order = list(K.CAT)
    rows.sort(key=lambda r: (r["ci"], order.index(r["id"]), {A: 0, P: 1, O: 2}[r["statut"]]))
    return rows


CONSIGNES = {
    "msp1": [
        "Chaîne d'alimentation : panneau 6 V → TP4056 → 18650 → boost réglé à 5,1 V → J1 (bornier).",
        "JP1 ôté (parqué en 2-3) : capteurs, OLED et servos coupés pendant le sommeil.",
        "JP12 en 2-3 (pont 1S 100k/100k) ; JP16-JP19 FERMÉS (LDR tracker) ; JP21 FERMÉ (DHT extérieur) ; JP13-JP15 et JP20 OUVERTS.",
        "NE PAS poser : LED5/R13, R17-R19, R27, tout l'étage relais (K1-K6…), Hi-Link, section 12 V.",
        "Firmware : seuil de veille d'urgence à régler pour ≥ 3,3 V cellule. Cellule à l'ombre.",
    ],
    "n3pp": [
        "Chargeur 5 V 2 A → cordon court et épais → J2 (jack). Jamais par le micro-USB du DevKit.",
        "JP1 ôté (parqué en 2-3) ; JP12 en 2-3 ; JP20 FERMÉ (LDR) ; JP16-JP19 OUVERTS (sondes de sol) ; JP5 sans cavalier (AUTO).",
        "Un seul relais : K1 = pompe d'arrosage (et ses pièces Q1, D1, R1, R5, LED1, R9, J3, JP5, R61).",
        "NE PAS poser : LED5/R13, R43-R46, R17-R19, K2-K6, servos, Hi-Link, section 12 V.",
        "J25 : relier J25-1 au +5V (J18-1) par un fil, sinon la mesure « batterie » lit 0 et le firmware se met en veille d'urgence (ou clé serveur 112 = 0). À vérifier au premier allumage.",
    ],
    "ffp_230": [
        "230 V → J27 (1 = N, 2 = L) → F1 T1A + RV1 → Hi-Link PS1 → +5V. Boîtier isolant fermé, différentiel 30 mA, supervision adulte.",
        "H1 et H5 : vis NYLON obligatoires. JP1 FERMÉ 1-2 (rail permanent) ; JP13-JP15 FERMÉS (ponts écho) ; JP20 FERMÉ ; JP5-JP10 sans cavalier (AUTO).",
        "J1 seulement pour les essais sur table SANS secteur : jamais J1 et le Hi-Link en même temps.",
        "Flash USB : carte HORS secteur et JP1 ôté (le pull-up 1-Wire sur GPIO2 bloque le bootloader). Sinon OTA.",
        "NE PAS poser : J29-J34, R43-R47, section 12 V (J26, Q11, D8, D9, R40, C5, J36, J37).",
    ],
    "ffp_12v": [
        "Batterie AGM (bornes batterie, pas la sortie « charge » du PWM) → fusible lame 7,5-10 A → J26 → Q11/D8/D9 → J36 → buck réglé 5,1 V → J37.",
        "JP1 FERMÉ ; JP12 en 1-2 (18k) ; JP13-JP15 FERMÉS ; JP20 FERMÉ ; JP5-JP10 sans cavalier (AUTO).",
        "Les pompes 12 V passent par les CONTACTS des relais, directement sur le bus ; diode de roue libre aux bornes de chaque pompe.",
        "Relais collés = poste dominant (≈ 10 Wh/jour par pompe maintenue) : cycles de pompe et chauffage interdit sur batterie (firmware).",
        "NE PAS poser : Hi-Link PS1, F1, RV1, J27, J29-J34, R43-R47.",
    ],
}


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def header(ws, row, names):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font, c.fill = f_head, fill_head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30


def fill_for(statut, source=None):
    if statut == O:
        return fill_opt
    if statut == P:
        return fill_own
    if source == "Import LCSC":
        return fill_imp
    return None


def build(OUT):
    wb = Workbook()

    # ================================================================== feuille Prix
    wp = wb.active
    wp.title = "Prix"
    wp["A1"] = "Catalogue des prix — fournisseurs marocains"
    wp["A1"].font = f_title
    wp["A2"] = (f"Relevé le {DATE} sur les pages produits, prix en DH tels qu'affichés (TTC). Cellules bleues = à mettre à jour ; "
                "toutes les autres feuilles se recalculent depuis celle-ci.")
    wp["A2"].font = f_sub
    PH = ["ID", "Article", "Fournisseur", "Prix du lot (DH)", "Taille du lot", "Unité de vente", "Prix unitaire (DH)",
          "Source", "Lien", "Remarque achat", "Relevé le"]
    header(wp, 4, PH)
    ids = list(K.CAT) + [s[0] for s in K.SHARED]
    names = {**{k: v[1] for k, v in K.CAT.items()}, **{s[0]: s[1] for s in K.SHARED}}
    for i, aid in enumerate(ids):
        r = 5 + i
        four, lot, n, unite, url, src, rem = PRIX[aid]
        vals = [aid, names[aid], four, lot, n, unite, f'=IF(D{r}="","",D{r}/E{r})', src, "lien" if url else "", rem, DATE]
        for j, v in enumerate(vals, 1):
            c = wp.cell(row=r, column=j, value=v)
            c.font = f_base
            c.alignment = wrap if j in (2, 10) else top
            c.border = border
        for j in (4, 5):
            wp.cell(row=r, column=j).font = f_input
        wp.cell(row=r, column=4).number_format = MONEY
        wp.cell(row=r, column=7).number_format = MONEY
        if url:
            wp.cell(row=r, column=9).hyperlink = url
            wp.cell(row=r, column=9).font = f_link
        if src == "Import LCSC":
            for j in range(1, 12):
                wp.cell(row=r, column=j).fill = fill_imp
    PLAST = 4 + len(ids)
    widths(wp, [16, 44, 20, 12, 9, 16, 12, 12, 7, 60, 11])
    wp.freeze_panes = "C5"
    wp.auto_filter.ref = f"A4:K{PLAST}"
    PR = f"Prix!$A$5:$A${PLAST}"


    def look(col, r):
        return f"INDEX(Prix!${col}$5:${col}${PLAST},MATCH($A{r},{PR},0))"


    # ================================================================== feuilles projet
    PCOLS = ["ID", "Catégorie", "Article", "Réf. carte", "Rôle", "Qté", "Statut", "Fournisseur", "Prix unitaire (DH)",
             "Coût (DH)", "Remarques", "Lien"]
    proj_total_cells = {}
    proj_rows = {}
    for prof, title in I.PROFILES.items():
        ws = wb.create_sheet(I.SHEET[prof])
        rows = profile_rows(prof)
        proj_rows[prof] = rows
        ws["A1"] = title
        ws["A1"].font = f_title
        ws["A2"] = (f"Carte n3-universal rev 0.2.2 — composants à poser et à acheter pour UNE unité. Prix du {DATE} "
                    "(feuille Prix). Les quantités (bleu) sont modifiables ; l'achat groupé arrondi aux lots est dans « Global ».")
        ws["A2"].font = f_sub
        r = 4
        ws.cell(row=r, column=1, value="Consignes de pose de ce profil").font = f_bold
        for line in CONSIGNES[prof]:
            r += 1
            c = ws.cell(row=r, column=2, value="• " + line)
            c.font = f_base
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 27 if len(line) > 130 else 15
        r += 2
        sum_row = r
        hdr = sum_row + 6
        first, last = hdr + 1, hdr + len(rows)
        J = f"J{first}:J{last}"
        G = f"G{first}:G{last}"
        I_ = f"I{first}:I{last}"
        labels = [
            ("Total à acheter, une unité (DH)", f'=SUMIFS({J},{G},"{A}")'),
            ("Options (DH)", f'=SUMIFS({J},{G},"{O}")'),
            ("Déjà possédé — valeur indicative (DH)", f'=SUMIFS({J},{G},"{P}")'),
            ("Articles à acheter sans prix en ligne (import LCSC, quincaillerie, PCB)",
             f'=SUMPRODUCT(({G}="{A}")*({I_}=""))'),
        ]
        for k, (lab, fml) in enumerate(labels):
            rr = sum_row + k
            a = ws.cell(row=rr, column=3, value=lab)
            a.font = f_bold
            a.fill = fill_tot
            v = ws.cell(row=rr, column=10, value=fml)
            v.font = f_bold
            v.fill = fill_tot
            v.number_format = MONEY if k < 3 else "0"
            for j in range(4, 10):
                ws.cell(row=rr, column=j).fill = fill_tot
        proj_total_cells[prof] = (f"J{sum_row}", f"J{sum_row + 1}")
        header(ws, hdr, PCOLS)
        prev, band = None, 0
        for i, row in enumerate(rows):
            rr = first + i
            if row["cat"] != prev:
                band ^= 1
                prev = row["cat"]
            src = PRIX[row["id"]][5]
            rem = row["remarque"]
            vals = [row["id"], row["cat"], row["article"], row["refs"], row["role"], row["qty"], row["statut"],
                    f"={look('C', rr)}", f'=IFERROR({look("G", rr)},"")', f'=IF(I{rr}="","",F{rr}*I{rr})', rem,
                    "lien" if PRIX[row["id"]][4] else ""]
            fl = fill_for(row["statut"], src) or fill_cat[band]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=rr, column=j, value=v)
                c.font = f_base
                c.alignment = wrap if j in (3, 4, 5, 11) else top
                c.border = border
                c.fill = fl
            ws.cell(row=rr, column=1).font = Font(name=F, size=8, color="888888")
            ws.cell(row=rr, column=6).font = f_input
            ws.cell(row=rr, column=8).font = f_link_sheet
            ws.cell(row=rr, column=9).font = f_link_sheet
            ws.cell(row=rr, column=9).number_format = MONEY
            ws.cell(row=rr, column=10).number_format = MONEY
            if PRIX[row["id"]][4]:
                ws.cell(row=rr, column=12).hyperlink = PRIX[row["id"]][4]
                ws.cell(row=rr, column=12).font = f_link
            if rem and any(w in rem for w in ("NE PAS", "OBLIGATOIRE", "JAMAIS", "ATTENTION", "INTERDIT")):
                ws.cell(row=rr, column=11).font = f_warn
        widths(ws, [13, 22, 34, 22, 40, 6, 12, 15, 11, 11, 58, 6])
        ws.freeze_panes = ws.cell(row=hdr + 1, column=4)
        ws.auto_filter.ref = f"A{hdr}:L{last}"
        ws.sheet_view.zoomScale = 90

    # ================================================================== feuille Global
    wg = wb.create_sheet("Global", 0)
    wg["A1"] = "Liste globale d'achats — 4 unités n3-universal (msp1, n3pp, ffp5cs 230 V, ffp5cs 12 V)"
    wg["A1"].font = f_title
    wg["A2"] = ("Quantités tirées des feuilles projet (formules) + rechange conseillée (bleu, modifiable), arrondies aux lots "
                f"vendus par les boutiques. Prix du {DATE}, feuille Prix.")
    wg["A2"].font = f_sub
    GCOLS = ["ID", "Catégorie", "Article", "Statut", "msp1", "n3pp", "ffp5cs 230V", "ffp5cs 12V", "Total 4 unités",
             "Rechange", "Qté à acheter", "Taille du lot", "Nb de lots", "Prix du lot (DH)", "Coût (DH)", "Fournisseur",
             "Source", "Remarque achat", "Lien"]
    # lignes globales : (id, statut) présents dans au moins un profil
    keys = []
    for prof in I.PROFILES:
        for row in proj_rows[prof]:
            k = (row["id"], row["statut"])
            if k not in keys:
                keys.append(k)
    order = list(K.CAT)
    keys.sort(key=lambda k: (K.CAT[k[0]][0], order.index(k[0]), {A: 0, P: 1, O: 2}[k[1]]))
    sum_row = 4
    hdr = 14
    first = hdr + 1
    last = hdr + len(keys)
    shared_hdr = last + 3
    shared_first = shared_hdr + 1
    shared_last = shared_first + len(K.SHARED) - 1
    O_ = f"O{first}:O{last}"
    D_ = f"D{first}:D{last}"
    N_ = f"N{first}:N{last}"
    K_ = f"K{first}:K{last}"
    glabels = [
        ("Total à acheter, 4 unités, arrondi aux lots (DH)", f'=SUMIFS({O_},{D_},"{A}")+SUM(O{shared_first}:O{shared_last})'),
        ("  dont outillage et consommables partagés (DH)", f"=SUM(O{shared_first}:O{shared_last})"),
        ("Options, si toutes prises (DH)", f'=SUMIFS({O_},{D_},"{O}")'),
        ("Déjà possédé — valeur indicative (DH)", f'=SUMIFS({O_},{D_},"{P}")'),
        ("Lignes à acheter sans prix en ligne (import LCSC, quincaillerie, PCB)",
         f'=SUMPRODUCT(({D_}="{A}")*({K_}>0)*({N_}=""))'),
    ]
    for k, (lab, fml) in enumerate(glabels):
        rr = sum_row + k
        a = wg.cell(row=rr, column=3, value=lab)
        a.font = f_bold
        for j in range(3, 16):
            wg.cell(row=rr, column=j).fill = fill_tot
        v = wg.cell(row=rr, column=15, value=fml)
        v.font = f_bold
        v.number_format = MONEY if k < 4 else "0"
    # coût par unité (somme des feuilles projet, sans arrondi aux lots)
    rr = sum_row + 6
    wg.cell(row=rr, column=3, value="Coût d'UNE unité (à l'unité, sans arrondi aux lots)").font = f_bold
    for j, prof in enumerate(I.PROFILES):
        wg.cell(row=rr, column=5 + j, value=I.SHEET[prof]).font = f_bold
        c = wg.cell(row=rr + 1, column=5 + j, value=f"='{I.SHEET[prof]}'!{proj_total_cells[prof][0]}")
        c.font = f_link_sheet
        c.number_format = MONEY
        c2 = wg.cell(row=rr + 2, column=5 + j, value=f"='{I.SHEET[prof]}'!{proj_total_cells[prof][1]}")
        c2.font = f_link_sheet
        c2.number_format = MONEY
    wg.cell(row=rr + 1, column=3, value="À acheter (DH)").font = f_base
    wg.cell(row=rr + 2, column=3, value="Options (DH)").font = f_base

    header(wg, hdr, GCOLS)
    prev, band = None, 0
    for i, (aid, st) in enumerate(keys):
        r = first + i
        ci, des, spare = K.CAT[aid]
        if K.CATS[ci] != prev:
            band ^= 1
            prev = K.CATS[ci]
        qf = []
        for prof in I.PROFILES:
            s = I.SHEET[prof]
            qf.append(f"=SUMIFS('{s}'!$F:$F,'{s}'!$A:$A,$A{r},'{s}'!$G:$G,$D{r})")
        if st != A:
            spare = 0
        vals = [aid, K.CATS[ci], des, st, *qf, f"=SUM(E{r}:H{r})", spare,
                f"=IF(I{r}=0,0,I{r}+J{r})", f"={look('E', r)}", f"=IF(K{r}=0,0,ROUNDUP(K{r}/L{r},0))",
                f'=IF({look("D", r)}="","",{look("D", r)})', f'=IF(N{r}="","",M{r}*N{r})', f"={look('C', r)}",
                f"={look('H', r)}", PRIX[aid][6], "lien" if PRIX[aid][4] else ""]
        fl = fill_for(st, PRIX[aid][5]) or fill_cat[band]
        for j, v in enumerate(vals, 1):
            c = wg.cell(row=r, column=j, value=v)
            c.font = f_base
            c.alignment = wrap if j in (3, 18) else top
            c.border = border
            c.fill = fl
        wg.cell(row=r, column=1).font = Font(name=F, size=8, color="888888")
        for j in range(5, 9):
            wg.cell(row=r, column=j).font = f_link_sheet
        wg.cell(row=r, column=9).font = f_bold
        wg.cell(row=r, column=10).font = f_input
        wg.cell(row=r, column=14).number_format = MONEY
        wg.cell(row=r, column=15).number_format = MONEY
        if PRIX[aid][4]:
            wg.cell(row=r, column=19).hyperlink = PRIX[aid][4]
            wg.cell(row=r, column=19).font = f_link

    wg.cell(row=shared_hdr - 1, column=1, value="Outillage et consommables partagés (une fois pour les 4 unités)").font = f_bold
    header(wg, shared_hdr, GCOLS)
    for i, (aid, des, rem) in enumerate(K.SHARED):
        r = shared_first + i
        vals = [aid, "Outillage partagé", des, A, "", "", "", "", 4 if aid == "DUPONT" else 1, 0, f"=I{r}+J{r}", f"={look('E', r)}",
                f"=ROUNDUP(K{r}/L{r},0)", f'=IF({look("D", r)}="","",{look("D", r)})', f'=IF(N{r}="","",M{r}*N{r})',
                f"={look('C', r)}", f"={look('H', r)}", rem + " " + PRIX[aid][6], "lien" if PRIX[aid][4] else ""]
        for j, v in enumerate(vals, 1):
            c = wg.cell(row=r, column=j, value=v)
            c.font = f_base
            c.alignment = wrap if j in (3, 18) else top
            c.border = border
        wg.cell(row=r, column=9).font = f_input
        wg.cell(row=r, column=1).font = Font(name=F, size=8, color="888888")
        wg.cell(row=r, column=14).number_format = MONEY
        wg.cell(row=r, column=15).number_format = MONEY
        if PRIX[aid][4]:
            wg.cell(row=r, column=19).hyperlink = PRIX[aid][4]
            wg.cell(row=r, column=19).font = f_link
    widths(wg, [13, 22, 38, 12, 10, 10, 11, 11, 9, 9, 9, 7, 7, 10, 11, 15, 11, 52, 6])
    wg.freeze_panes = wg.cell(row=hdr + 1, column=4)
    wg.auto_filter.ref = f"A{hdr}:S{last}"
    wg.sheet_view.zoomScale = 90

    # ================================================================== Lisez-moi
    wl = wb.create_sheet("Lisez-moi", 0)
    L = [
        ("Achats n3-universal — 4 projets", f_title),
        (f"Carte universelle rev 0.2.2 (hardware/n3-universal). Prix relevés le {DATE} chez les boutiques marocaines en ligne.", f_sub),
        ("", None),
        ("Feuilles", f_bold),
        ("Global — tout ce qu'il faut acheter pour les 4 unités, rechange comprise, arrondi aux lots vendus, avec le coût total.", f_base),
        ("msp1, n3pp, ffp5cs 230V, ffp5cs 12V — une feuille par projet : consignes de pose, composants à poser sur la carte et ce qui se branche autour, coût d'une unité.", f_base),
        ("Prix — catalogue unique (fournisseur, prix, lien). Changer un prix ici met tout le classeur à jour.", f_base),
        ("", None),
        ("Cellules modifiables (texte bleu)", f_bold),
        ("Qté dans les feuilles projet ; Rechange dans Global ; Prix du lot et Taille du lot dans Prix. Le reste est calculé.", f_base),
        ("", None),
        ("Statuts et couleurs", f_bold),
        ("À acheter — compté dans les totaux. Fond rose = introuvable au Maroc, à mettre dans le colis LCSC combiné avec la commande JLCPCB (COMMANDE.md §2).", f_base),
        ("Option — fond jaune : utile mais pas indispensable (BME280, INA226 supplémentaires, microSD). Total séparé.", f_base),
        ("Déjà possédé — fond vert : TP4056, batterie AGM, régulateur PWM, panneau 12 V. Valeur indicative, hors total.", f_base),
        ("", None),
        ("Points à retenir avant d'acheter", f_bold),
        ("1. Trois pièces ne se trouvent pas au Maroc : les deux TVS (1.5KE6.8A, 1.5KE18A) → LCSC ; prévoir aussi le PCB (JLCPCB). Les ajouter au même colis.", f_warn),
        ("2. Boîtiers : la carte mesure 278 × 135 mm → coffret d'au moins 300 × 200 mm.", f_base),
        ("3. Borniers KF301 vendus « pas 5,0 mm » : ils entrent dans les trous 5,08 en 2 et 3 pôles ; en 4 pôles, présenter avant de souder.", f_base),
        ("4. Modules I2C (OLED, DS3231, BME280, INA226) : vérifier l'ordre des broches avant d'enficher (supports = GND-VCC-SCL-SDA), sinon fils Dupont.", f_warn),
        ("5. Lecteur microSD « pour Arduino » : type Catalex (régulateur + tampon) = ne convient pas. Chercher un module 3,3 V direct.", f_base),
        ("6. Fusible T1A : demander explicitement un temporisé (T). Porte-fusible : mesurer l'entraxe des pattes (22,0–22,6 mm).", f_base),
        ("7. Pas de panneau 6 V 5 W au Maroc : 6 V 3,8 W retenu (marge ×3 en hiver selon ALIMENTATION.md).", f_base),
        ("8. Presse-étoupes, porte-fusible lame et fil unipolaire : pas en ligne, quincaillerie / électricien / accessoires auto.", f_base),
        ("9. Règle absolue (D5 en pont de fil) : ne JAMAIS brancher l'USB du DevKit quand la carte est alimentée. Flasher carte hors tension, ou en OTA.", f_warn),
        ("", None),
        ("Mise à jour", f_bold),
        ("Classeur généré depuis la BOM : `python3 achats/build_achats.py` (dossier hardware/n3-universal). Prix : achats/prix.py ; "
         "pose par profil : achats/inventaire.py. La CI échoue si la BOM change sans que cette liste suive.", f_base),
        ("", None),
        ("Limites", f_bold),
        ("Prix tels qu'affichés au relevé ; stocks variables (plusieurs articles en rupture, signalés). Certains prix viennent d'une fourchette de variantes "
         "(bas de fourchette retenu) : à confirmer au panier. Les quantités de câble et de presse-étoupes sont des estimations à ajuster à l'installation.", f_base),
    ]
    for i, (t, fnt) in enumerate(L, 1):
        c = wl.cell(row=i, column=1, value=t)
        if fnt:
            c.font = fnt
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wl.column_dimensions["A"].width = 140
    for i, (t, _) in enumerate(L, 1):
        if len(t) > 150:
            wl.row_dimensions[i].height = 30

    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.gridLines = False
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    return len(keys)
